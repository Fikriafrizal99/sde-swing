#!/usr/bin/env python3
from __future__ import annotations

"""Build FTJ Performance Setup from existing performance artifacts.

This remains a presentation/export layer only. Canonical trading outcomes are
not recalculated or changed.
"""

import argparse
import sys
from datetime import datetime
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

import pandas as pd
from openpyxl import load_workbook

# When this file is executed directly (``python tools/export_performance_workbook.py``),
# Python puts the ``tools`` directory on sys.path, not the repository root. Add the
# project root explicitly so the package-style imports below work both from the
# maintenance BAT menu and when imported by tests/other modules.
PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from tools.performance_report.analytics import (
    broker_snapshot,
    generated_frames,
    read_csv,
    setup_frames,
)
from tools.performance_report.dashboard import BRAND_NAME, build_dashboard
from tools.performance_report.styling import style_sheet

DEFAULT_INPUT = PROJECT_ROOT / "data/output/analytics/performance"
DEFAULT_OUTPUT = DEFAULT_INPUT / "exports"
OUTPUT_PREFIX = "FTJ_PERFORMANCE_SETUP"

REPORT_SPECS = (
    ("PERFORMANCE_SUMMARY.csv", "Performance Summary", True),
    ("PERFORMANCE_BY_SETUP.csv", "Setup Performance", False),
    ("PERFORMANCE_BY_SIGNAL_TYPE.csv", "Signal Performance", False),
    ("PERFORMANCE_BY_BROKER_CONFIDENCE.csv", "Broker Confidence", False),
    ("PERFORMANCE_BY_MARKET_REGIME.csv", "Market Regime", False),
    ("BROKER_PERIOD_PERFORMANCE.csv", "Broker Period", False),
    ("BROKER_CONFIDENCE_BY_PERIOD.csv", "Broker Conf x Period", False),
    ("EXIT_EFFICIENCY_SUMMARY.csv", "Exit Summary", False),
    ("EXIT_EFFICIENCY_BY_SETUP.csv", "Exit by Setup", False),
    ("EXIT_EFFICIENCY_TRADES.csv", "Exit Trades", False),
    ("EXIT_EFFICIENCY_INTEGRITY.csv", "Integrity Flags", False),
    ("SIGNAL_OUTCOME_LEDGER.csv", "Signal Ledger", False),
    ("ACTIVE_RECOMMENDATIONS.csv", "Active Recommendations", False),
    ("LIFECYCLE_EVENTS.csv", "Lifecycle Events", False),
    ("SIGNAL_RECOMMENDATION_HISTORY.csv", "Recommendation History", False),
    ("PORTFOLIO_POSITIONS.csv", "Actual Portfolio", False),
)


def _validate_written_workbook(path: Path) -> None:
    """Fail fast if the produced XLSX is not structurally readable.

    This does not recalculate any performance. It only validates that the
    Office Open XML archive is intact and can be reopened by openpyxl after the
    writer has fully closed it.
    """
    try:
        with ZipFile(path, "r") as archive:
            broken = archive.testzip()
            if broken:
                raise RuntimeError(f"XLSX archive rusak pada part: {broken}")
            names = set(archive.namelist())
            required_parts = {
                "[Content_Types].xml",
                "xl/workbook.xml",
                "xl/styles.xml",
                "xl/worksheets/sheet1.xml",
            }
            missing_parts = sorted(required_parts - names)
            if missing_parts:
                raise RuntimeError(f"XLSX kehilangan part wajib: {', '.join(missing_parts)}")

            # Excel-safe report intentionally does not emit Structured Table
            # definitions. The sheets remain filterable/formatted ranges.
            table_parts = [name for name in names if name.startswith("xl/tables/")]
            if table_parts:
                raise RuntimeError(
                    "XLSX masih mengandung Structured Table definition yang tidak diharapkan: "
                    + ", ".join(sorted(table_parts)[:5])
                )
    except BadZipFile as exc:
        raise RuntimeError(f"File hasil export bukan XLSX/ZIP yang valid: {path}") from exc

    probe = load_workbook(path, read_only=False, data_only=False)
    try:
        if "Dashboard" not in probe.sheetnames:
            raise RuntimeError("Sheet Dashboard tidak ditemukan setelah workbook dibuka ulang.")
        if probe["Dashboard"]["A1"].value != BRAND_NAME:
            raise RuntimeError("Header FTJ Performance Setup tidak valid setelah workbook dibuka ulang.")
    finally:
        probe.close()


def build_workbook(
    input_dir: Path,
    output_dir: Path,
    *,
    output_path: Path | None = None,
) -> tuple[Path, pd.DataFrame]:
    input_dir = input_dir.expanduser().resolve()
    output_dir = output_dir.expanduser().resolve()
    core = input_dir / "PERFORMANCE_SUMMARY.csv"
    if not core.exists():
        raise FileNotFoundError(
            f"Core performance belum tersedia: {core}. Jalankan Update outcome terlebih dahulu."
        )

    output_dir.mkdir(parents=True, exist_ok=True)
    if output_path is None:
        output_path = output_dir / f"{OUTPUT_PREFIX}_{datetime.now().astimezone():%Y%m%d_%H%M%S}.xlsx"
    else:
        output_path = output_path.expanduser().resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)

    summary = read_csv(core)
    ledger = read_csv(input_dir / "SIGNAL_OUTCOME_LEDGER.csv")
    portfolio = read_csv(input_dir / "PORTFOLIO_POSITIONS.csv")
    integrity = read_csv(input_dir / "EXIT_EFFICIENCY_INTEGRITY.csv")
    exit_summary = read_csv(input_dir / "EXIT_EFFICIENCY_SUMMARY.csv")
    setup_full, setup_snap = setup_frames(read_csv(input_dir / "PERFORMANCE_BY_SETUP.csv"))
    broker_snap = broker_snapshot(read_csv(input_dir / "PERFORMANCE_BY_BROKER_CONFIDENCE.csv"))
    generated = generated_frames(
        summary,
        ledger,
        portfolio,
        integrity,
        exit_summary,
        setup_snap,
        broker_snap,
    )

    manifest_rows: list[dict[str, Any]] = []
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        pd.DataFrame().to_excel(writer, sheet_name="Dashboard", index=False, header=False)

        for sheet, frame in generated.items():
            frame.to_excel(writer, sheet_name=sheet, index=False)
            manifest_rows.append(
                {
                    "Source_File": "GENERATED_PRESENTATION_LAYER",
                    "Sheet": sheet,
                    "Status": "INCLUDED",
                    "Rows": len(frame),
                }
            )

        for filename, sheet, required in REPORT_SPECS:
            path = input_dir / filename
            if not path.exists():
                manifest_rows.append(
                    {
                        "Source_File": filename,
                        "Sheet": sheet,
                        "Status": "MISSING_REQUIRED" if required else "MISSING_OPTIONAL",
                        "Rows": 0,
                    }
                )
                if required:
                    raise FileNotFoundError(f"Required performance artifact tidak ditemukan: {path}")
                continue

            if filename == "PERFORMANCE_SUMMARY.csv":
                frame = summary
            elif filename == "PERFORMANCE_BY_SETUP.csv":
                frame = setup_full
            else:
                frame = read_csv(path)
            frame.to_excel(writer, sheet_name=sheet, index=False)
            manifest_rows.append(
                {"Source_File": filename, "Sheet": sheet, "Status": "INCLUDED", "Rows": len(frame)}
            )

        manifest = pd.DataFrame(manifest_rows)
        manifest.to_excel(writer, sheet_name="Manifest", index=False)
        for sheet, worksheet in writer.sheets.items():
            if sheet != "Dashboard":
                style_sheet(worksheet, table=(sheet != "Overview"))
        build_dashboard(writer, summary)

    _validate_written_workbook(output_path)
    return output_path, pd.DataFrame(manifest_rows)


def make_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=f"Export {BRAND_NAME} ke workbook Excel dashboard multi-sheet"
    )
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--output-file", type=Path, default=None)
    return parser


def main() -> int:
    args = make_parser().parse_args()
    path, manifest = build_workbook(args.input_dir, args.output_dir, output_path=args.output_file)
    included = int((manifest["Status"] == "INCLUDED").sum()) if not manifest.empty else 0
    missing = int((manifest["Status"] == "MISSING_OPTIONAL").sum()) if not manifest.empty else 0
    print(f"{BRAND_NAME} berhasil dibuat dan lolos validasi XLSX.")
    print(f"Included sections : {included}")
    print(f"Optional missing  : {missing}")
    print(f"File              : {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
