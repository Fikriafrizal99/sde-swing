from __future__ import annotations

from datetime import datetime
from typing import Any

import pandas as pd
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .analytics import metric, num
from .styling import NAVY, SLATE, WHITE, status_fill

BRAND_NAME = "FTJ Performance Setup"
MUTED, BORDER = "64748B", "CBD5E1"


def _column(ws, name: str) -> int | None:
    for cell in ws[1]:
        if str(cell.value or "") == name:
            return cell.column
    return None


def _bar(ws, category: str, value: str, title: str, horizontal: bool = True) -> BarChart | None:
    ccol, vcol = _column(ws, category), _column(ws, value)
    if ws.max_row < 2 or not ccol or not vcol:
        return None
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.style = 10
    chart.height, chart.width, chart.title = 7.2, 14.2, title
    if horizontal:
        chart.y_axis.title, chart.x_axis.title = category.replace("_", " "), value.replace("_", " ")
    else:
        chart.x_axis.title, chart.y_axis.title = category.replace("_", " "), value.replace("_", " ")
    chart.add_data(Reference(ws, min_col=vcol, min_row=1, max_row=ws.max_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=ccol, min_row=2, max_row=ws.max_row))
    chart.legend = None
    return chart


def _status(summary: pd.DataFrame) -> tuple[str, str]:
    if summary.empty:
        return "NO DATA", "Core performance summary kosong."
    row = summary.iloc[0]
    closed = int(num(metric(row, "Closed", "Closed_Outcomes")))
    wr, avg, pf = num(metric(row, "Win_Rate_Pct")), num(metric(row, "Average_Return_Pct")), num(metric(row, "Profit_Factor"))
    if closed < 20:
        return "LOW SAMPLE", "Edge belum cukup sampel untuk dinilai stabil."
    if avg > 0 and pf >= 1.3 and wr >= 50:
        if closed < 100:
            return "PROMISING EDGE", "Positif, tetapi masih butuh sampel dan market regime tambahan."
        return "POSITIVE EDGE", "Canonical sample sudah lebih matang; tetap monitor regime dan integrity."
    if avg < 0 or (pf and pf < 1):
        return "NEEDS REVIEW", "Canonical expectancy belum sehat pada periode report."
    return "WATCH", "Belum ada sinyal statistik yang cukup kuat untuk klasifikasi lebih tinggi."


def _card(ws, start: int, end: int, row: int, label: str, value: Any, fmt: str) -> None:
    ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    ws.merge_cells(start_row=row + 1, start_column=start, end_row=row + 2, end_column=end)
    title, data = ws.cell(row, start), ws.cell(row + 1, start)
    title.value, data.value = label, value
    title.fill, data.fill = PatternFill("solid", fgColor=SLATE), PatternFill("solid", fgColor="F8FAFC")
    title.font, data.font = Font(color=WHITE, bold=True, size=10), Font(color=NAVY, bold=True, size=18)
    title.alignment = data.alignment = Alignment(horizontal="center", vertical="center")
    data.number_format = fmt
    thin = Side(style="thin", color=BORDER)
    for r in range(row, row + 3):
        for c in range(start, end + 1):
            ws.cell(r, c).border = Border(left=thin, right=thin, top=thin, bottom=thin)


def build_dashboard(writer: pd.ExcelWriter, summary: pd.DataFrame) -> None:
    ws = writer.sheets["Dashboard"]
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    for col in range(1, 15):
        ws.column_dimensions[get_column_letter(col)].width = 12

    ws.merge_cells("A1:N2")
    ws["A1"] = BRAND_NAME
    ws["A1"].font = Font(size=24, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")

    row = summary.iloc[0] if not summary.empty else pd.Series(dtype=object)
    ws.merge_cells("A3:N3")
    ws["A3"] = (f"Canonical KPI dashboard | Periode {metric(row, 'Start_Date', default='-')} s.d. "
                  f"{metric(row, 'End_Date', default='-')} | Generated {datetime.now().astimezone():%Y-%m-%d %H:%M %Z}")
    ws["A3"].font = Font(color=MUTED, italic=True, size=10)

    cards = [
        ("Closed", metric(row, "Closed", "Closed_Outcomes"), "0"),
        ("Win Rate", metric(row, "Win_Rate_Pct"), '0.00"%"'),
        ("Avg Return", metric(row, "Average_Return_Pct"), '0.00"%"'),
        ("Profit Factor", metric(row, "Profit_Factor"), "0.00"),
        ("Expectancy", metric(row, "Expectancy_Pct"), '0.00"%"'),
        ("Trigger Rate", metric(row, "Trigger_Rate_Pct"), '0.00"%"'),
        ("Avg Hold", metric(row, "Average_Holding_Days"), '0.00" d"'),
        ("Excluded", metric(row, "Canonical_Excluded_Episodes"), "0"),
    ]
    spans = [(1, 3), (4, 6), (7, 9), (10, 12)]
    for idx, (label, value, fmt) in enumerate(cards):
        _card(ws, *spans[idx % 4], 5 if idx < 4 else 9, label, value, fmt)

    status, note = _status(summary)
    ws.merge_cells("A13:N13")
    ws["A13"] = f"REPORT STATUS: {status}"
    ws["A13"].font = Font(size=14, bold=True, color=NAVY)
    ws["A13"].fill = PatternFill("solid", fgColor=status_fill(status))
    ws["A13"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A14:N14")
    ws["A14"] = note
    ws["A14"].font = Font(color=SLATE, italic=True)
    ws["A14"].alignment = Alignment(horizontal="center")

    ws["P1"], ws["Q1"] = "Outcome", "Count"
    for r, name in enumerate(("Win", "Loss", "Ambiguous"), 2):
        ws.cell(r, 16, name)
        ws.cell(r, 17, metric(row, name))
    ws.column_dimensions["P"].hidden = ws.column_dimensions["Q"].hidden = True
    pie = PieChart()
    pie.title, pie.style, pie.height, pie.width = "Canonical closed outcomes", 10, 7.2, 12
    pie.add_data(Reference(ws, min_col=17, min_row=1, max_row=4), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=16, min_row=2, max_row=4))
    ws.add_chart(pie, "A16")

    if "Setup Snapshot" in writer.sheets:
        chart = _bar(writer.sheets["Setup Snapshot"], "Setup", "Average_Return_Pct", "Average return by setup", False)
        if chart:
            ws.add_chart(chart, "H16")
        chart = _bar(writer.sheets["Setup Snapshot"], "Setup", "Win_Rate_Pct", "Win rate by setup", False)
        if chart:
            ws.add_chart(chart, "A31")
    if "Score Analysis" in writer.sheets:
        chart = _bar(writer.sheets["Score Analysis"], "Score_Bucket", "Average_Return_Pct", "Exploratory return by score bucket")
        if chart:
            ws.add_chart(chart, "H31")

    eq = writer.sheets.get("Equity Curve")
    if eq is not None and eq.max_row >= 3 and _column(eq, "Exit_Date") and _column(eq, "Exploratory_Index"):
        chart = LineChart()
        chart.title, chart.style, chart.height, chart.width = "Exploratory raw-ledger equity curve (index 100)", 13, 7, 28.8
        chart.y_axis.title, chart.x_axis.title = "Index", "Exit sequence"
        chart.add_data(Reference(eq, min_col=_column(eq, "Exploratory_Index"), min_row=1, max_row=eq.max_row), titles_from_data=True)
        chart.set_categories(Reference(eq, min_col=_column(eq, "Exit_Date"), min_row=2, max_row=eq.max_row))
        chart.legend = None
        ws.add_chart(chart, "A46")

    ws.merge_cells("A61:N62")
    ws["A61"] = ("Catatan: KPI utama berasal dari canonical PERFORMANCE_SUMMARY. Score Analysis, Time Analysis, dan Equity Curve "
                  "berasal dari raw closed ledger dan bersifat exploratory; bukan pengganti canonical performance atau actual portfolio P&L.")
    ws["A61"].font = Font(color=MUTED, italic=True, size=9)
    ws["A61"].alignment = Alignment(wrap_text=True)
