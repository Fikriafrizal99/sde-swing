from __future__ import annotations

from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY, SLATE, WHITE = "0F172A", "334155", "FFFFFF"
GREEN_LIGHT, RED_LIGHT, AMBER_LIGHT, BLUE_LIGHT = "DCFCE7", "FEE2E2", "FEF3C7", "DBEAFE"
GRID = "E2E8F0"


def status_fill(text: str) -> str:
    text = text.upper()
    if any(x in text for x in ("STRONG", "PROMISING", "POSITIVE", "OK", "SOURCE OF TRUTH")):
        return GREEN_LIGHT
    if any(x in text for x in ("WEAK", "REVIEW", "NEEDS")):
        return RED_LIGHT
    if any(x in text for x in ("LOW SAMPLE", "WATCH")):
        return AMBER_LIGHT
    return BLUE_LIGHT


def style_sheet(ws, table: bool = True) -> None:
    """Style a worksheet without creating Excel Structured Table objects.

    The ``table`` argument is retained for API compatibility with the exporter,
    but deliberately no ``openpyxl.worksheet.table.Table`` is created. Raw SDE
    artifacts can contain complex/legacy headers that openpyxl is willing to
    serialize while desktop Excel may reject as a structured-table definition.
    AutoFilter + header styling provide the same day-to-day usability without
    adding that compatibility risk.
    """
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    if ws.max_row < 1 or ws.max_column < 1:
        return

    # Keep the range filterable like a table, without emitting xl/tables/*.xml.
    if ws.max_row >= 2:
        ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=SLATE)
    ws.row_dimensions[1].height = 28

    thin = Side(style="thin", color=GRID)
    for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 5000)):
        for cell in row:
            cell.border = Border(bottom=thin)

    headers = {str(c.value or ""): c.column for c in ws[1]}
    for header, col in headers.items():
        letter = get_column_letter(col)
        upper = header.upper()
        if any(x in upper for x in ("PCT", "RATE", "RETURN", "MFE", "MAE", "EFFICIENCY", "CONFIDENCE")):
            for cell in ws[letter][1:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '0.00"%"'
        if any(x in upper for x in ("RETURN_PCT", "EXPECTANCY_PCT", "WIN_RATE_PCT", "PROFIT_FACTOR")) and ws.max_row >= 2:
            ws.conditional_formatting.add(
                f"{letter}2:{letter}{ws.max_row}",
                ColorScaleRule(
                    start_type="min",
                    start_color="FECACA",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FEF3C7",
                    end_type="max",
                    end_color="BBF7D0",
                ),
            )

    for name in ("Report_Status", "Status", "Sample_Status"):
        if name in headers:
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row, headers[name])
                cell.fill = PatternFill("solid", fgColor=status_fill(str(cell.value or "")))
                cell.font = Font(bold=True, color=NAVY)

    for col in ws.iter_cols():
        letter = col[0].column_letter
        width = max((len(str(c.value or "")) for c in col[: min(len(col), 250)]), default=10) + 2
        ws.column_dimensions[letter].width = min(max(width, 11), 36)
